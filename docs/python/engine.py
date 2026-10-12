"""Tablova: independent, offline PDF to editable Excel conversion engine."""
from __future__ import annotations

import csv
import io
import os
import re
import statistics
import subprocess
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from threading import Event
from typing import Callable

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .table_geometry import find_ruled_tables, is_header, clean_page, is_watermark_char
from .document_layout import is_document_layout, collect_document_page, write_document_workbook

MAX_PAGES = 3000
MAX_BYTES = 512 * 1024 * 1024
MAX_ROWS = 1_048_576
MAX_COLS = 16_384
Progress = Callable[[int, int, str], None]


class ConversionError(Exception):
    """A conversion problem suitable for displaying to the end user."""


class ConversionCancelled(ConversionError):
    pass


@dataclass
class TableData:
    page: int
    rows: list[list[str]]
    method: str
    bbox: tuple[float, float, float, float] | None = None


@dataclass
class ConversionResult:
    path: Path
    pages: int
    tables: int
    rows: int
    ocr_pages: int
    warnings: list[str] = field(default_factory=list)
    watermark_pages: int = 0


def _check_cancel(stop: Event | None) -> None:
    if stop is not None and stop.is_set():
        raise ConversionCancelled("Dönüştürme kullanıcı tarafından iptal edildi.")


def _clean(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"[ \t]+", " ", str(value).replace("\x00", "").replace("\r", "")).strip()


def _sanitize(rows: list[list[object]]) -> list[list[str]]:
    cleaned = [[_clean(v) for v in row] for row in rows if row]
    cleaned = [r for r in cleaned if any(r)]
    if not cleaned:
        return []
    width = max(len(r) for r in cleaned)
    if width > MAX_COLS:
        return []
    return [r + [""] * (width - len(r)) for r in cleaned]


def _score(rows: list[list[str]]) -> float:
    if len(rows) < 2 or len(rows[0]) < 2:
        return -1.0
    count = sum(bool(c) for r in rows for c in r)
    density = count / max(1, len(rows) * len(rows[0]))
    multi = sum(sum(bool(c) for c in r) >= 2 for r in rows)
    if multi < 2 or density < 0.30:
        return -1.0
    return density * 10 + min(6, multi) * 0.7 + min(len(rows[0]), 12) * 0.1


def _iou(a: tuple, b: tuple) -> float:
    x0, y0 = max(a[0], b[0]), max(a[1], b[1])
    x1, y1 = min(a[2], b[2]), min(a[3], b[3])
    inter = max(0.0, x1 - x0) * max(0.0, y1 - y0)
    if not inter:
        return 0.0
    area_a = max(0, (a[2] - a[0]) * (a[3] - a[1]))
    area_b = max(0, (b[2] - b[0]) * (b[3] - b[1]))
    return inter / max(1.0, area_a + area_b - inter)


def _native_tables(page, page_number: int) -> list[TableData]:
    # Strong evidence from actual ruled cells wins over text heuristics.
    ruled, overlay = find_ruled_tables(page)
    if ruled:
        return [TableData(page_number, _sanitize(item.rows), "digital/ruled",
                          item.bbox) for item in ruled if sum(bool(x) for r in item.rows for x in r) >= 3]
    # Text-only PDFs still need a fallback. Filter diagonal overlays first.
    clean, _ = clean_page(page)
    try:
        candidates = clean.find_tables(table_settings={
            "vertical_strategy": "text", "horizontal_strategy": "text",
            "min_words_vertical": 2, "min_words_horizontal": 1})
        selected = []
        for item in candidates:
            rows = _sanitize(item.extract())
            if _score(rows) >= 8.0 and len(rows) >= 3:
                selected.append(TableData(page_number, rows, "digital/text", item.bbox))
        return selected
    finally:
        if clean is not page:
            clean.flush_cache()

def _runtime_paths() -> tuple[Path, Path]:
    import sys
    root = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent.parent))
    base = root / "runtime" / "tesseract"
    exe = base / "bin" / "tesseract.exe"
    tess = base / "tessdata"
    return exe, tess


def _ocr_words(pdf_path: Path, page_index: int, *, stop: Event | None) -> list[dict]:
    _check_cancel(stop)
    exe, tess = _runtime_paths()
    if not exe.is_file() or not (tess / "tur.traineddata").is_file():
        raise ConversionError("Taranmış belge için yerel metin tanıma bileşeni bulunamadı.")
    import pypdfium2 as pdfium
    source = pdfium.PdfDocument(str(pdf_path))
    try:
        pdf_page = source[page_index]
        scale = 300 / 72
        try:
            image = pdf_page.render(scale=scale).to_pil().convert("RGB")
        finally:
            pdf_page.close()
    finally:
        source.close()
    with tempfile.TemporaryDirectory(prefix="tablova_ocr_") as tempdir:
        bitmap = Path(tempdir) / "page.png"
        image.save(bitmap)
        image.close()
        cmd = [str(exe), str(bitmap), "stdout", "--tessdata-dir",
               str(tess), "-l", "tur+eng", "--psm", "6", "tsv"]
        startup = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        try:
            run = subprocess.run(cmd, capture_output=True, timeout=120,
                                 startupinfo=None, creationflags=startup)
        except subprocess.TimeoutExpired as exc:
            raise ConversionError(f"OCR işlem süresi aşıldı (sayfa {page_index + 1}).") from exc
        except OSError as exc:
            raise ConversionError(f"OCR çalıştırılamadı: {exc}") from exc
        _check_cancel(stop)
        if run.returncode:
            raise ConversionError(
                f"OCR sayfa {page_index + 1} işlenemedi (kod {run.returncode})."
            )
        parsed = csv.DictReader(io.StringIO(run.stdout.decode("utf-8", errors="replace")),
                                delimiter="\t")
        words = []
        for word in parsed:
            try:
                conf = float(word.get("conf", "-1"))
                text = _clean(word.get("text"))
                if not text or conf < 12:
                    continue
                x, y = int(word["left"]), int(word["top"])
                width, height = int(word["width"]), int(word["height"])
                words.append({"text": text, "x": x, "y": y,
                              "w": width, "h": height,
                              "block": word.get("block_num", ""),
                              "line": word.get("line_num", ""),
                              "par": word.get("par_num", "")})
            except (ValueError, KeyError):
                continue
        return words


def _ocr_table(words: list[dict], page_number: int) -> TableData | None:
    """Conservative OCR line segmentation; never guesses absent text."""
    if len(words) < 4:
        return None
    grouped: dict[tuple, list[dict]] = {}
    for word in words:
        key = (word["block"], word["par"], word["line"])
        grouped.setdefault(key, []).append(word)
    lines = sorted(grouped.values(), key=lambda ws: (min(w["y"] for w in ws),
                                                     min(w["x"] for w in ws)))
    if len(lines) < 2:
        return None
    segments: list[list[tuple[float, str]]] = []
    for line in lines:
        ordered = sorted(line, key=lambda w: w["x"])
        h = statistics.median([w["h"] for w in ordered])
        threshold = max(25.0, h * 2.3)
        row: list[tuple[float, str]] = []
        group: list[dict] = []
        right = 0
        for word in ordered:
            if group and word["x"] - right > threshold:
                row.append((float(group[0]["x"]), " ".join(w["text"] for w in group)))
                group = []
            group.append(word)
            right = max(right, word["x"] + word["w"])
        if group:
            row.append((float(group[0]["x"]), " ".join(w["text"] for w in group)))
        segments.append(row)
    # Require repeated multi-column structure instead of inventing a table.
    counts = Counter(len(row) for row in segments if len(row) >= 2)
    if not counts:
        return None
    columns, repeats = counts.most_common(1)[0]
    if columns > 30 or repeats < 2:
        return None
    reference = [r for r in segments if len(r) == columns]
    anchors = [statistics.median(r[i][0] for r in reference) for i in range(columns)]
    if any(anchors[i + 1] - anchors[i] < 16 for i in range(columns - 1)):
        return None
    built: list[list[str]] = []
    for segment_row in segments:
        if len(segment_row) < 2:
            continue
        row = [""] * columns
        for x, value in segment_row:
            k = min(range(columns), key=lambda i: abs(x - anchors[i]))
            row[k] = (row[k] + " " + value).strip() if row[k] else value
        built.append(row)
    final = _sanitize(built)
    if _score(final) < 0:
        return None
    return TableData(page_number, final, "ocr/lines")


def _header_equal(a: list[str], b: list[str]) -> bool:
    if len(a) != len(b) or len(a) < 2:
        return False
    normalize = lambda s: re.sub(r"\s+", " ", s.lower()).strip()
    return [normalize(v) for v in a] == [normalize(v) for v in b] and sum(bool(x) for x in a) >= 2


def _join_continuations(tables: list[TableData]) -> list[TableData]:
    results: list[TableData] = []
    for item in tables:
        if not item.rows:
            continue
        prev = results[-1] if results else None
        if prev and prev.page + 1 == item.page and len(prev.rows[0]) == len(item.rows[0]):
            # An explicitly repeated header is not another transaction.
            if _header_equal(prev.rows[0], item.rows[0]) or (is_header(prev.rows[0]) and is_header(item.rows[0])):
                prev.rows.extend(item.rows[1:])
                prev.page = item.page
                continue
            # PDFs commonly print the 12-column report heading only on page 1.
            # Adjacent ruled grids of the same schema continue one logical table.
            if (len(item.rows[0]) == 12 and prev.method == "digital/ruled"
                    and item.method == "digital/ruled" and prev.bbox and item.bbox
                    and abs(prev.bbox[0] - item.bbox[0]) <= 10
                    and not is_header(item.rows[0])):
                prev.rows.extend(item.rows)
                prev.page = item.page
                continue
        results.append(item)
    return results

_NUMBER = re.compile(r"^-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?$")
_SIMPLE = re.compile(r"^-?\d+(?:\.\d+)?$")
_DATE = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")


def _typed(raw: str) -> tuple[object, str]:
    """Conservative Turkish numbers; ambiguous IDs stay strings."""
    s = raw.strip()
    if not s:
        return None, "empty"
    if _DATE.fullmatch(s):
        try:
            return datetime.strptime(s, "%d.%m.%Y").date(), "date"
        except ValueError:
            return s, "text"
    negative_parens = s.startswith("(") and s.endswith(")")
    sample = s[1:-1].strip() if negative_parens else s
    sample = re.sub(r"(?i)(?:₺|TL|TRY)\s*", "", sample).strip()
    sample = sample.replace("\u00a0", "").replace(" ", "")
    if sample.endswith("%"):
        return s, "text"
    if _NUMBER.fullmatch(sample) or _SIMPLE.fullmatch(sample):
        digits = re.sub(r"\D", "", sample)
        if len(digits) > 15 or (len(digits) > 1 and digits[0] == "0"
                                 and not sample.startswith("0,")):
            return s, "text"
        if re.match(r"^-?\d+\.\d{1,2}$", sample) and "," not in sample:
            # Dot fractional values in a Turkish document can be ambiguous.
            return s, "text"
        normalized = sample.replace(".", "").replace(",", ".") if "," in sample or "." in sample and _NUMBER.fullmatch(sample) else sample
        try:
            numeric = float(normalized)
            return -abs(numeric) if negative_parens else numeric, "number"
        except ValueError:
            pass
    return s, "text"


def _set_cell(cell, raw: str) -> None:
    if cell.column in (5, 6, 7) and "\n" in str(raw):
        joined = re.sub(r"\s+", "", str(raw))
        if (cell.column == 5 and re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", joined)) or (
                cell.column in (6, 7) and re.fullmatch(
                    r"-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d{1,2})?", joined)):
            raw = joined
    value, kind = _typed(raw)
    if kind == "empty":
        return
    if kind == "date":
        cell.value = value
        cell.number_format = "DD.MM.YYYY"
        return
    if kind == "number":
        cell.value = value
        cell.number_format = "#,##0.00" if "," in raw else "#,##0.########"
        return
    # openpyxl interprets a leading '=' as an executable formula.
    cell.value = value
    cell.data_type = "s"


def _sheet_name(index: int) -> str:
    return f"Tablo_{index}"


# Readable, bounded widths for the 12-column report. Other documents use
# safe adaptive widths rather than a single 46-character column cap.
REPORT_WIDTHS = (17, 25, 24, 39, 18, 18, 19, 15, 27, 67, 73, 35)


def _write_workbook(tables: list[TableData], output: Path, stop: Event | None) -> int:
    import math
    wb = Workbook()
    wb.remove(wb.active)
    written = 0
    for tab in tables:
        _check_cancel(stop)
        if not tab.rows:
            continue
        if len(tab.rows[0]) == 12 and is_header(tab.rows[0]):
            tab.rows[0] = [
                "Takip No", "Ad Soyad", "Hizmet Türü", "İşlem",
                "İşlem Tarihi", "İşlem Tutarı", "Kesinti Tutarı",
                "Grup Türü", "Grup Adı", "Kesinti Nedeni",
                "Değerlendirme", "Sonuç"
            ]
        default_widths = list(REPORT_WIDTHS) if len(tab.rows[0]) == 12 else [
            20 if i < 5 else 32 for i in range(len(tab.rows[0]))]
        current_sheet = None
        for index, data in enumerate(tab.rows):
            if index % 200 == 0:
                _check_cancel(stop)
            if index % MAX_ROWS == 0:
                current_sheet = wb.create_sheet(_sheet_name(len(wb.sheetnames) + 1))
                current_sheet.freeze_panes = "A2"
                current_sheet.sheet_view.showGridLines = False
                for col, width in enumerate(default_widths, 1):
                    current_sheet.column_dimensions[get_column_letter(col)].width = width
            sheet_row = (index % MAX_ROWS) + 1
            if current_sheet is None:
                raise ConversionError("Excel çalışma sayfası oluşturulamadı.")
            needed_lines = 1
            for col, raw in enumerate(data[:MAX_COLS], 1):
                cell = current_sheet.cell(sheet_row, col)
                _set_cell(cell, raw)
                if col in (6, 7) and "#" in str(raw):
                    cell.fill = PatternFill("solid", fgColor="FFF1C2")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                col_width = default_widths[col - 1] if col <= len(default_widths) else 24
                lines = str(raw).split("\n")
                lines_needed = sum(max(1, math.ceil(len(line) / max(8, col_width - 4)))
                                   for line in lines)
                needed_lines = max(needed_lines, lines_needed)
                if sheet_row == 1:
                    cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1A3550")
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
            # Excel maximum row height is 409.5pt; avoid hidden text by
            # accommodating typical wrapped cells and capping pathological ones.
            current_sheet.row_dimensions[sheet_row].height = min(405, max(
                26 if sheet_row == 1 else 19, needed_lines * 15 + 7))
            written += 1
        if current_sheet is not None and current_sheet.max_row > 1:
            current_sheet.auto_filter.ref = current_sheet.dimensions
    if not wb.worksheets:
        raise ConversionError("Excel'e aktarılabilecek tablo bulunamadı.")
    wb.save(output)
    return written

def _verify_workbook(path: Path, expected_tables: int) -> None:
    if not path.is_file() or path.stat().st_size <= 0:
        raise ConversionError("Excel dosyası oluşturulamadı.")
    probe = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(probe.sheetnames) < expected_tables:
            raise ConversionError("Excel doğrulaması başarısız: eksik sayfa.")
        if any(probe[s].max_row < 2 for s in probe.sheetnames):
            raise ConversionError("Excel doğrulaması başarısız: boş tablo.")
    finally:
        probe.close()


def _publish_no_clobber(staging: Path, wanted: Path) -> Path:
    for num in range(1, 10000):
        target = wanted if num == 1 else wanted.with_name(f"{wanted.stem} ({num}){wanted.suffix}")
        try:
            os.link(staging, target)
            return target
        except FileExistsError:
            continue
        except OSError:
            # Windows rename does not replace an existing destination.
            if os.name == "nt" and not target.exists():
                try:
                    staging.rename(target)
                    return target
                except FileExistsError:
                    continue
            raise ConversionError("Excel hedef klasöre güvenli şekilde kaydedilemedi.")
    raise ConversionError("Uygun Excel dosya adı üretilemedi.")


def convert_pdf(
    source: str | Path,
    output: str | Path | None = None,
    *,
    progress: Progress | None = None,
    stop: Event | None = None,
) -> ConversionResult:
    input_path = Path(source).expanduser().resolve()
    if not input_path.is_file():
        raise ConversionError("Seçilen PDF dosyası bulunamadı.")
    if not (0 < input_path.stat().st_size <= MAX_BYTES):
        raise ConversionError("PDF dosyası boş veya 512 MB sınırından büyük.")
    with input_path.open("rb") as f:
        if b"%PDF-" not in f.read(1024):
            raise ConversionError("Dosya geçerli PDF başlığı içermiyor.")
    wanted = Path(output).expanduser().resolve() if output else input_path.with_suffix(".xlsx")
    wanted.parent.mkdir(parents=True, exist_ok=True)
    if wanted.suffix.lower() != ".xlsx":
        raise ConversionError("Hedef dosya .xlsx uzantılı olmalı.")
    if progress:
        progress(0, 1, "PDF inceleniyor")
    raw_tables: list[TableData] = []
    document_pages = []
    warnings: list[str] = []
    ocr_pages = 0
    watermark_pages = 0
    _check_cancel(stop)
    try:
        with pdfplumber.open(str(input_path)) as pdf:
            pages = len(pdf.pages)
            if pages == 0 or pages > MAX_PAGES:
                raise ConversionError("PDF sayfa sayısı desteklenen aralığın dışında.")
            document_mode = is_document_layout(pdf)
            for pi, page in enumerate(pdf.pages):
                _check_cancel(stop)
                if any(is_watermark_char(c) for c in page.chars):
                    watermark_pages += 1
                try:
                    if document_mode:
                        document_pages.append(collect_document_page(page, pi + 1))
                        if progress:
                            progress(pi + 1, pages + 1, f"Belge sayfası {pi + 1}/{pages} okundu")
                        continue
                    try:
                        results = _native_tables(page, pi + 1) if len(page.chars) > 4 else []
                    except Exception as exc:
                        results = []
                        warnings.append(f"Sayfa {pi + 1}: dijital tablo analizi başarısız ({type(exc).__name__}).")
                    if not results and len(page.chars) < 12:
                        try:
                            words = _ocr_words(input_path, pi, stop=stop)
                            ocr_pages += 1
                            table = _ocr_table(words, pi + 1)
                            if table is not None:
                                results.append(table)
                        except ConversionCancelled:
                            raise
                        except ConversionError as exc:
                            warnings.append(f"Sayfa {pi + 1}: {exc}")
                    if not results:
                        warnings.append(f"Sayfa {pi + 1}: güvenilir tablo bulunamadı.")
                    raw_tables.extend(results)
                finally:
                    page.flush_cache()
                if progress:
                    progress(pi + 1, pages + 1, f"Sayfa {pi + 1}/{pages} işlendi")
    except ConversionCancelled:
        raise
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"PDF okunamadı: {type(exc).__name__}: {exc}") from exc
    masked_amount_cells = sum(
        "#" in str(value)
        for tab in raw_tables for record in tab.rows
        for value in record[5:7]
    )
    if masked_amount_cells:
        warnings.append(
            f"Kaynak PDF'de {masked_amount_cells} tutar hücresi okunabilir sayı yerine "
            "###### içeriyor; bu hücreler Excel'de sarı işaretlendi."
        )
    tables = _join_continuations(raw_tables) if not document_pages else []
    if not tables and not document_pages:
        raise ConversionError("PDF içinde güvenilir şekilde çıkarılabilen tablo bulunamadı. " +
                              (" ".join(warnings[:2]) if warnings else ""))
    _check_cancel(stop)
    staging: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(prefix=".tablova-", suffix=".xlsx",
                                         dir=str(wanted.parent), delete=False) as temp:
            staging = Path(temp.name)
        if progress:
            progress(pages, pages + 1, "Excel oluşturuluyor ve doğrulanıyor")
        row_count = (write_document_workbook(document_pages, staging, stop)
                     if document_pages else _write_workbook(tables, staging, stop))
        _check_cancel(stop)
        _verify_workbook(staging, 1 if document_pages else len(tables))
        _check_cancel(stop)
        final = _publish_no_clobber(staging, wanted)
        if progress:
            progress(pages + 1, pages + 1, "Tamamlandı")
        if document_pages:
            warnings.append("Belge düzeni: metinler sayfa sırasıyla, çizgili tablolar ayrı Excel sayfalarında aktarıldı. QR ve görseller Excel'e aktarılmaz.")
        return ConversionResult(final, pages, len(tables) + sum(len(p.tables) for p in document_pages), row_count, ocr_pages, warnings, watermark_pages)
    finally:
        if staging is not None:
            staging.unlink(missing_ok=True)
