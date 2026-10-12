"""Loss-aware editable export for short PDF documents with mixed layouts.

No OCR, cloud calls, images, hidden heuristic replacements or business-specific
name/identity parsing. Full extracted text lines are authoritative.
"""
from __future__ import annotations

import math
from dataclasses import dataclass
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass
class DocumentPage:
    number: int
    lines: list[tuple[float, float, str]]
    tables: list[list[list[str]]]


def _ruled_tables(page):
    return [
        t for t in page.find_tables(table_settings={
            "vertical_strategy": "lines", "horizontal_strategy": "lines"
        })
        if len(t.rows) >= 2 and len(t.columns) >= 2
    ]


def is_document_layout(pdf) -> bool:
    """Route short official forms and mixed tabular documents, not bulk tables."""
    if not (1 <= len(pdf.pages) <= 3):
        return False
    for page in pdf.pages:
        words = page.extract_words()
        if len(words) < 5:
            continue
        tables = _ruled_tables(page)
        if not tables or len(tables) > 1:
            return True
        inside = 0
        for w in words:
            x = (w["x0"] + w["x1"]) / 2
            y = (w["top"] + w["bottom"]) / 2
            if any(t.bbox[0] <= x <= t.bbox[2]
                   and t.bbox[1] <= y <= t.bbox[3] for t in tables):
                inside += 1
        if inside / len(words) < 0.84:
            return True
    return False


def collect_document_page(page, number: int) -> DocumentPage:
    lines = []
    for line in page.extract_text_lines():
        text = str(line.get("text") or "").replace("\x00", "").strip()
        if not text:
            continue
        # Preserve original text and line order; do not infer IDs, dates or
        # numeric values because leading zeros and separators are significant.
        while text:
            part, text = text[:32000], text[32000:]
            lines.append((round(float(line["top"]), 2),
                          round(float(line["x0"]), 2), part))
    tables = []
    for table in _ruled_tables(page):
        rows = [
            [str(value).strip() if value is not None else ""
             for value in row]
            for row in table.extract()
        ]
        if len(rows) >= 2 and any(any(v for v in row) for row in rows):
            tables.append(rows)
    return DocumentPage(number, lines, tables)


def _title(sheet, columns: Sequence[str]):
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="153550")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 29
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False


def _text(cell, value):
    # Explicitly store as string, including content starting with "=".
    cell.value = str(value)
    cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_document_workbook(pages: list[DocumentPage], target, stop=None) -> int:
    wb = Workbook()
    wb.remove(wb.active)
    total = 0
    for p in pages:
        if stop is not None and stop.is_set():
            raise RuntimeError("Dönüştürme kullanıcı tarafından iptal edildi.")
        sh = wb.create_sheet(f"Belge_{p.number}")
        _title(sh, ("Satır", "Dikey konum", "Yatay konum", "Belgedeki metin"))
        for col, width in (("A", 9), ("B", 15), ("C", 15), ("D", 105)):
            sh.column_dimensions[col].width = width
        if not p.lines:
            raise ValueError(f"Sayfa {p.number} için seçilebilir metin bulunamadı.")
        for i, (top, x0, text) in enumerate(p.lines, 2):
            sh.cell(i, 1, i-1)
            sh.cell(i, 2, top)
            sh.cell(i, 3, x0)
            _text(sh.cell(i, 4), text)
            sh.row_dimensions[i].height = min(105, max(18, math.ceil(len(text)/115)*16+4))
            total += 1
        sh.auto_filter.ref = sh.dimensions
        # Actual ruled tables are provided on their own editable worksheets.
        for rows in p.tables:
            if len(wb.worksheets) >= 100:
                raise ValueError("Belgede desteklenenden fazla bağımsız tablo var.")
            t = wb.create_sheet(f"Tablo_{len([s for s in wb if s.title.startswith('Tablo_')])+1}")
            width = max(map(len, rows))
            _title(t, [f"Sütun {i}" for i in range(1, width+1)])
            for c in range(1, width+1):
                t.column_dimensions[get_column_letter(c)].width = 27
            for ri, row in enumerate(rows, 2):
                for ci, value in enumerate(row, 1):
                    _text(t.cell(ri,ci),value)
            t.freeze_panes = "A2"
            t.auto_filter.ref = t.dimensions
    if total == 0:
        raise ValueError("Excel'e aktarılacak metin bulunamadı.")
    wb.save(target)
    return total
